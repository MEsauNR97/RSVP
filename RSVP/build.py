"""
build.py — Genera index.html con los datos de familias.xlsx incrustados.

Uso:
    python build.py

Requiere: openpyxl
    pip install openpyxl
"""

import json
import re
import openpyxl

EXCEL_FILE = "familias.xlsx"
HTML_FILE  = "index.html"

# ── 1. Leer Excel ─────────────────────────────────────────────────────────────
wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
ws = wb.active

headers = [str(cell.value).strip() for cell in ws[1]]

familias = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    data = dict(zip(headers, row))
    codigo  = str(data.get("codigo", "")).strip().zfill(3)
    familia = str(data.get("familia", "")).strip()
    adultos = int(data.get("adultos") or 0)
    ninos   = int(data.get("ninos")   or 0)
    extra   = int(data.get("extra")   or 0)

    if codigo and familia:
        familias[codigo] = {
            "familia": familia,
            "adultos": adultos,
            "ninos":   ninos,
            "extra":   extra,
        }

print(f"  OK: {len(familias)} familias leidas de {EXCEL_FILE}")

# ── 2. Generar bloque JS ──────────────────────────────────────────────────────
lines = ["    const familias = {"]
for codigo, datos in familias.items():
    lines.append(
        f"      '{codigo}': "
        f"{{ familia: {json.dumps(datos['familia'], ensure_ascii=False)}, "
        f"adultos: {datos['adultos']}, "
        f"ninos: {datos['ninos']}, "
        f"extra: {datos['extra']} }},"
    )
lines.append("    };")
nuevo_bloque = "\n".join(lines)

# ── 3. Reemplazar bloque en index.html ────────────────────────────────────────
with open(HTML_FILE, "r", encoding="utf-8") as f:
    html = f.read()

patron = re.compile(
    r"// ─── BASE DE DATOS DE FAMILIAS.*?// ───────────────────────────────────────────────────────────",
    re.DOTALL
)

reemplazo = (
    "// ─── BASE DE DATOS DE FAMILIAS ──────────────────────────────────────────\n"
    "    // Generado automáticamente por build.py — no editar a mano\n"
    + nuevo_bloque + "\n"
    "    // ───────────────────────────────────────────────────────────"
)

html_nuevo, n = patron.subn(reemplazo, html)

if n == 0:
    print("  ERROR: No se encontro el bloque de familias en index.html.")
    print("    Asegurate de que el archivo tenga los marcadores correctos.")
    exit(1)

with open(HTML_FILE, "w", encoding="utf-8") as f:
    f.write(html_nuevo)

print(f"  OK: {HTML_FILE} actualizado con {len(familias)} familias.")
print("  Listo — abre index.html en el navegador.")
